"""业务知识库（RAG）：业务/<名>/知识/ 下的 md/txt/csv/json/xlsx/docx → 切块 → 关键词(FTS5 BM25) + 向量 两路召回 → RRF 融合。
agent 调 search_kb 查，结果带「文件#位置」出处。文件有改动下次查询时自动重建，不用手动。

向量：默认本地 TF-IDF 字符 n-gram + SVD（不联网、不花钱）；设 EMBED_BASE_URL/EMBED_API_KEY/EMBED_MODEL 就走 OpenAI 兼容的 /embeddings。
"""
import json
import os
import re
import sqlite3

import numpy as np

import 环境

目录 = 环境.业务 / "知识"
库 = 环境.记忆 / "知识.db"
_向量文件 = 环境.记忆 / "知识向量.npz"
读法 = {".md", ".txt", ".csv", ".json", ".xlsx", ".docx"}
_缓存 = {}


def _读文件(p) -> list[tuple[str, str]]:
    """→ [(位置, 文本)]。表格按行成段，带表头，模型才看得懂每个数是什么。"""
    if p.suffix == ".xlsx":
        from openpyxl import load_workbook
        段 = []
        for ws in load_workbook(p, read_only=True, data_only=True).worksheets:
            行 = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
            if not 行:
                continue
            头 = 行[0]
            for i in range(1, len(行), 20):
                块 = "\n".join("；".join(f"{h}={v}" for h, v in zip(头, r) if v) for r in 行[i:i + 20])
                段.append((f"{ws.title}!{i + 1}", 块))
        return 段
    if p.suffix == ".docx":
        import docx
        文 = "\n".join(x.text for x in docx.Document(p).paragraphs)
    else:
        文 = p.read_text(encoding="utf-8", errors="ignore")
    return [(f"段{i + 1}", 块) for i, 块 in enumerate(_切(文))]


def _切(文: str, 长=500, 叠=80) -> list[str]:
    """按段落攒到 ~500 字一块，块间重叠 80 字，别把一条规则切两半。"""
    块, 当前 = [], ""
    for 段 in re.split(r"\n\s*\n", 文):
        if len(当前) + len(段) > 长 and 当前:
            块.append(当前)
            当前 = 当前[-叠:]
        当前 += ("\n\n" if 当前 else "") + 段
    return 块 + ([当前] if 当前.strip() else [])


def _指纹() -> str:
    fs = sorted(p for p in 目录.rglob("*") if p.suffix in 读法) if 目录.exists() else []
    return json.dumps([(str(p.relative_to(目录)), p.stat().st_mtime, p.stat().st_size) for p in fs])


class _连接(sqlite3.Connection):
    def __exit__(self, *a):          # with 结束提交并关闭；不关的话 Windows 上文件一直被锁
        super().__exit__(*a)
        self.close()


def _连():
    c = sqlite3.connect(库, factory=_连接)
    c.executescript("create table if not exists 块(id integer primary key, 文件 text, 位置 text, 文本 text);"
                    "create virtual table if not exists 块索引 using fts5(文本, content='块', content_rowid='id', tokenize='trigram');"
                    "create table if not exists 元(k text primary key, v text);")
    return c


class _本地向量:
    def 训练(self, 文本):
        from sklearn.decomposition import TruncatedSVD
        from sklearn.feature_extraction.text import TfidfVectorizer
        self.tf = TfidfVectorizer(analyzer="char_wb", ngram_range=(2, 3), sublinear_tf=True, max_features=20000)
        x = self.tf.fit_transform(文本)
        self.svd = TruncatedSVD(max(1, min(128, x.shape[1] - 1, len(文本) - 1)), random_state=0).fit(x) if len(文本) > 2 else None
        return self

    def 编码(self, 文本):
        x = self.tf.transform(文本)
        v = self.svd.transform(x) if self.svd else x.toarray()
        return (v / (np.linalg.norm(v, axis=1, keepdims=True) + 1e-9)).astype(np.float32)


class _接口向量:
    def 训练(self, 文本):
        return self

    def 编码(self, 文本):
        import httpx
        out = []
        for i in range(0, len(文本), 64):
            r = httpx.post(os.environ["EMBED_BASE_URL"].rstrip("/") + "/embeddings", timeout=60,
                           headers={"Authorization": f"Bearer {os.environ.get('EMBED_API_KEY', '')}"},
                           json={"model": os.environ.get("EMBED_MODEL", "text-embedding-3-small"), "input": 文本[i:i + 64]})
            r.raise_for_status()
            out += [d["embedding"] for d in r.json()["data"]]
        v = np.array(out, dtype=np.float32)
        return v / (np.linalg.norm(v, axis=1, keepdims=True) + 1e-9)


def 重建() -> int:
    段 = [(str(p.relative_to(目录)), 位, 文) for p in sorted(目录.rglob("*")) if p.suffix in 读法
          for 位, 文 in _读文件(p) if 文.strip()] if 目录.exists() else []
    with _连() as c:
        c.execute("delete from 块")
        c.executemany("insert into 块(文件, 位置, 文本) values(?,?,?)", 段)
        c.execute("insert into 块索引(块索引) values('rebuild')")
        c.execute("insert or replace into 元 values('指纹', ?)", (_指纹(),))
    _缓存.clear()
    if 段:
        器 = (_接口向量() if os.environ.get("EMBED_BASE_URL") else _本地向量()).训练([x[2] for x in 段])
        _缓存.update(器=器, 向量=器.编码([x[2] for x in 段]))
    return len(段)


def _确保最新():
    with _连() as c:
        旧 = c.execute("select v from 元 where k='指纹'").fetchone()
    if not 旧 or 旧[0] != _指纹() or "器" not in _缓存:
        重建()


def 查(问题: str, k: int = 5) -> list[dict]:
    _确保最新()
    if "器" not in _缓存:
        return []
    with _连() as c:
        ids = [r[0] for r in c.execute("select id from 块 order by id")]
        sims = _缓存["向量"] @ _缓存["器"].编码([问题])[0]
        排名 = [[ids[i] for i in np.argsort(-sims)[:30]]]
        词 = [w for w in re.findall(r"[a-zA-Z0-9_]+|[一-鿿]+", 问题)]
        长 = {g for w in 词 for g in ([w[i:i + 3] for i in range(len(w) - 2)] if len(w) >= 3 else [])}
        短 = [w for w in 词 if len(w) < 3]
        if 长:   # trigram 至少 3 字；短词走 LIKE
            排名.append([r[0] for r in c.execute("select rowid from 块索引 where 块索引 match ? order by bm25(块索引) limit 30",
                                               (" OR ".join('"' + w.replace('"', "") + '"' for w in 长),))])
        if 短:
            行 = c.execute("select id, 文本 from 块 where " + " or ".join(["文本 like ?"] * len(短)) + " limit 500",
                          [f"%{w}%" for w in 短]).fetchall()
            排名.append([i for i, t in sorted(行, key=lambda r: -sum(r[1].count(w) for w in 短))[:30]])
        分 = {}
        for 表 in 排名:
            for r, i in enumerate(表):
                分[i] = 分.get(i, 0) + 1 / (60 + r)
        top = sorted(分, key=lambda i: -分[i])[:k]
        行 = {r[0]: r for r in c.execute(f"select id, 文件, 位置, 文本 from 块 where id in ({','.join('?' * len(top))})", top)} if top else {}
    return [{"出处": f"{行[i][1]}#{行[i][2]}", "文本": 行[i][3], "分": round(分[i], 4)} for i in top]
